"""Writing cells into the workbook in place: formulas, their stored results and everything else survive a pull."""
import re
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

from labmap import model, xlsx

ROOT = Path(__file__).resolve().parent.parent
SHEET = "placeables"


def part(path, sheet=SHEET):
    with zipfile.ZipFile(path) as z:
        return xlsx._parts(z)[sheet]


def read_part(path, name=None):
    with zipfile.ZipFile(path) as z:
        return z.read(name or part(path)).decode("utf-8")


def ref_of(path, id, column, sheet=SHEET):
    """The cell holding one row's column, e.g. 'G12'."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path, read_only=True)
    try:
        ws = wb[sheet]
        head = {}
        for n, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if n == 1:
                head = {v.strip(): k for k, v in enumerate(row, start=1) if isinstance(v, str)}
            elif n > 2 and row[head["id"] - 1] == id:
                return f"{get_column_letter(head[column])}{n}"
        raise AssertionError(f"{id} isn't on the {sheet} sheet")
    finally:
        wb.close()


def put_raw(path, ref, xml, sheet=SHEET):
    """Drop a cell straight into the sheet, the way Excel writes one: a formula with its stored result."""
    name = part(path, sheet)
    with zipfile.ZipFile(path) as z:
        raw = {n: z.read(n) for n in z.namelist()}
        infos = z.infolist()
    sheet = raw[name].decode("utf-8")
    sheet, n = re.subn(r'<c\s+r="%s"(\s[^>]*?)?(/>|>.*?</c>)' % ref, xml, sheet, count=1, flags=re.S)
    assert n == 1, f"{ref} isn't in the sheet"
    raw[name] = sheet.encode("utf-8")
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for info in infos:
            out.writestr(info, raw[info.filename])
    shutil.move(str(tmp), str(path))


class InPlace(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        for name in ("lab-data.xlsx", "rooms", "shapes"):
            src = ROOT / "example" / name
            (shutil.copytree if src.is_dir() else shutil.copy2)(src, self.tmp / name)
        self.path = self.tmp / "lab-data.xlsx"
        self.x_of_spec = ref_of(self.path, "SPEC-02", "x")
        put_raw(self.path, self.x_of_spec, '<c r="%s"><f>3+7</f><v>10</v></c>' % self.x_of_spec)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def write(self, moves):
        return model.write_moves(self.path, moves, self.tmp / "backups")

    def test_the_fixture_reads_like_excel_wrote_it(self):
        self.assertEqual(model.load(self.tmp).placeables["SPEC-02"]["x"], 10)  # the stored result, not the formula

    def test_a_formula_elsewhere_keeps_its_stored_result(self):
        self.write({"placeables": {"CART-02": {"x": 123}}})
        self.assertIn("<f>3+7</f><v>10</v>", read_part(self.path))
        lab = model.load(self.tmp)
        self.assertEqual(lab.placeables["SPEC-02"]["x"], 10)  # ...so it still reads as 10, not as blank
        self.assertEqual(lab.placeables["CART-02"]["x"], 123)
        self.assertEqual(lab.stale, [])

    def test_everything_it_did_not_touch_is_byte_for_byte(self):
        before = {n: read_part(self.path, n) for n in zipfile.ZipFile(self.path).namelist()}
        self.write({"placeables": {"CART-02": {"x": 123}}})
        after = {n: read_part(self.path, n) for n in zipfile.ZipFile(self.path).namelist()}
        self.assertEqual(sorted(before), sorted(after))
        changed = {n for n in before if before[n] != after[n]}
        # the sheet, and workbook.xml if it didn't already ask Excel to recalculate: no styles, no validation,
        # no other sheet, no dropped conditional formatting
        self.assertIn(part(self.path), changed)
        self.assertFalse(changed - {part(self.path), "xl/workbook.xml"})
        self.assertIn('fullCalcOnLoad="1"', after["xl/workbook.xml"])

    def test_writing_over_a_formula_reports_it_and_asks_excel_to_recalculate(self):
        _, _, note = self.write({"placeables": {"SPEC-02": {"x": 42}}})
        self.assertEqual(note["replaced"], [("SPEC-02", "x", "=3+7")])
        self.assertIsNone(note["rewrote"])
        self.assertNotIn("<f>3+7</f>", read_part(self.path))
        self.assertEqual(model.load(self.tmp).placeables["SPEC-02"]["x"], 42)
        self.assertIn('fullCalcOnLoad="1"', read_part(self.path, "xl/workbook.xml"))

    def test_numbers_text_and_blanks(self):
        self.write({"placeables": {"CART-02": {"x": 12.5, "faces": "NE", "notes": None}}})
        r = model.load(self.tmp).placeables["CART-02"]
        self.assertEqual((r["x"], r["faces"], r.get("notes")), (12.5, "NE", None))

    def test_a_missing_cell_is_created_in_the_right_place(self):
        """A blank cell isn't in the file at all; the new one has to go in column order or Excel complains."""
        blank = ref_of(self.path, "CART-02", "notes")
        put_raw(self.path, blank, "")
        self.assertNotIn(f'r="{blank}"', read_part(self.path))
        self.write({"placeables": {"CART-02": {"notes": "moved for the audit"}}})
        row = re.search(r'<row r="%s".*?</row>' % re.search(r"\d+", blank).group(0), read_part(self.path), re.S).group(0)
        refs = [xlsx.column_index(m) for m in re.findall(r'<c\s+r="([A-Z]+\d+)"', row)]
        self.assertEqual(refs, sorted(refs))
        self.assertEqual(model.load(self.tmp).placeables["CART-02"]["notes"], "moved for the audit")

    def test_a_shared_formula_falls_back_to_rewriting_the_workbook(self):
        """Other cells point at a shared formula's master, so that one isn't ours to edit: openpyxl takes over."""
        put_raw(self.path, self.x_of_spec,
                '<c r="%s"><f t="shared" ref="%s" si="3">3+7</f><v>10</v></c>' % (self.x_of_spec, self.x_of_spec))
        _, _, note = self.write({"placeables": {"SPEC-02": {"x": 42}}})
        self.assertIn("shared formula", note["rewrote"])
        self.assertEqual(note["replaced"], [("SPEC-02", "x", "=3+7")])
        self.assertEqual(model.load(self.tmp).placeables["SPEC-02"]["x"], 42)

    def test_an_id_that_is_not_there(self):
        _, missing, _ = self.write({"placeables": {"NOPE-99": {"x": 1}}})
        self.assertEqual(missing, ["NOPE-99"])


class SyncLists(unittest.TestCase):
    """Bringing a workbook made earlier up to date with the blank template's lists, dropdown and all."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.path = self.tmp / "lab-data.xlsx"
        shutil.copy2(ROOT / "example" / "lab-data.xlsx", self.path)
        self.col, self.last = model.list_layout(self.path)["category"]
        self.dropped = model.read_workbook(self.path)[1]["category"][-1][0]
        put_raw(self.path, f"A{self.last}", "", sheet="lists")  # a workbook from before that category existed

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def sync(self):
        return model.sync_lists(self.path, ROOT / "example" / "lab-data.xlsx", self.tmp / "backups")

    def test_it_adds_what_is_missing(self):
        self.assertEqual(self.sync(), {"category": [self.dropped]})
        self.assertIn(self.dropped, [v for v, _ in model.read_workbook(self.path)[1]["category"]])

    def test_the_dropdown_reaches_the_new_value(self):
        self.sync()
        book = read_part(self.path, "xl/workbook.xml")
        ref = re.search(r'name="L_category"[^>]*>([^<]*)<', book).group(1)
        self.assertEqual(ref, f"lists!$A$2:$A${model.list_layout(self.path)['category'][1]}")

    def test_it_leaves_a_workbook_that_is_already_current_alone(self):
        self.sync()
        before = read_part(self.path, "xl/workbook.xml")
        self.assertEqual(self.sync(), {})
        self.assertEqual(read_part(self.path, "xl/workbook.xml"), before)

    def test_it_keeps_what_the_workbook_already_had(self):
        xlsx.write_cells(self.path, {"lists": {f"A{self.last}": "ours-only"}})  # a category only this lab uses
        self.sync()
        values = [v for v, _ in model.read_workbook(self.path)[1]["category"]]
        self.assertIn("ours-only", values)
        self.assertIn(self.dropped, values)


if __name__ == "__main__":
    unittest.main()
