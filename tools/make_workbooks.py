"""Build lab-data.xlsx (empty template) and example/lab-data.xlsx (worked example)."""
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import example_data as ex
import lists
import schema

FIRST = 3  # row 1 = column names, row 2 = hints, data from row 3
DARK, LIGHT, HINT = "1F4E79", "DDEBF7", "F2F2F2"
SHEET_ROWS = {s[0]: s[3] for s in schema.SHEETS}


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def red():
    return dict(fill=fill("FFC7CE"), font=Font(color="9C0006"))


def orange():
    return dict(fill=fill("FFEB9C"), font=Font(color="9C5700"))


def header(cell, text, dark):
    cell.value = text
    cell.font = Font(bold=True, color="FFFFFF" if dark else "1F1F1F")
    cell.fill = fill(DARK if dark else LIGHT)


def validation(kind, arg, L):
    if kind == "id":
        return DataValidation(type="custom", errorTitle="ID format",
                              formula1=f'AND(EXACT({L}{FIRST},UPPER({L}{FIRST})),ISERROR(SEARCH(" ",{L}{FIRST})))',
                              error="IDs use UPPERCASE letters, digits, - and . with no spaces.")
    if kind == "int":
        return DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", errorTitle="Whole number",
                              error="Whole numbers only (cm, W, kg): no decimals, no units.")
    if kind == "pct":
        return DataValidation(type="whole", operator="between", formula1="0", formula2="100",
                              errorTitle="Percent full", error="A whole number from 0 to 100.")
    if kind == "list":
        return DataValidation(type="list", formula1=f"L_{arg}",
                              error="Pick a value from the list, or add a new one on the lists sheet.")
    if kind == "listw":
        return DataValidation(type="list", formula1=f"L_{arg}", errorStyle="warning",
                              error="Not in the list. Keep it anyway? If it's a new standard value, "
                                    "add it to the lists sheet.")
    if kind == "ref":
        return DataValidation(type="list", formula1=f"{arg}!$A${FIRST}:$A${FIRST + SHEET_ROWS[arg] - 1}",
                              errorStyle="warning", error=f"This ID isn't on the {arg} sheet (yet).")
    return None


def write_sheet(wb, name, tab, freeze, nrows, cols, rows):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = tab
    last = FIRST + nrows - 1
    letters = {c["name"]: get_column_letter(i) for i, c in enumerate(cols, 1)}
    for i, c in enumerate(cols, 1):
        L = letters[c["name"]]
        h = ws.cell(1, i)
        header(h, c["name"], c["req"])
        h.alignment = Alignment(horizontal="center")
        if c["help"]:
            h.comment = Comment(c["help"], "lab-map", width=320, height=36 + 16 * (len(c["help"]) // 42 + 1))
        t = ws.cell(2, i, c["hint"] or None)
        t.font = Font(italic=True, color="595959", size=9)
        t.fill = fill(HINT)
        t.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[L].width = c["width"] or 12
        rng = f"{L}{FIRST}:{L}{last}"
        kind, _, arg = c["kind"].partition(":")
        dv = validation(kind, arg, L)
        if dv:
            dv.allow_blank = True
            dv.showErrorMessage = True
            dv.add(rng)
            ws.add_data_validation(dv)
        if kind == "id":
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'AND(${L}{FIRST}<>"",COUNTIF(${L}${FIRST}:${L}${last},${L}{FIRST})>1)'], **red()))
        elif kind == "ref":
            tlast = FIRST + SHEET_ROWS[arg] - 1
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'AND(${L}{FIRST}<>"",COUNTIF({arg}!$A${FIRST}:$A${tlast},${L}{FIRST})=0)'], **red()))
        elif kind in ("textfmt", "date"):
            fmt = "@" if kind == "textfmt" else "yyyy-mm-dd"
            for r in range(FIRST, last + 1):
                ws.cell(r, i).number_format = fmt
    if name == "placeables":  # a child's room must match its parent's room
        P, R = letters["parent"], letters["room"]
        ws.conditional_formatting.add(f"{R}{FIRST}:{R}{last}", FormulaRule(formula=[
            f'AND(${P}{FIRST}<>"",IFERROR(INDEX(${R}${FIRST}:${R}${last},'
            f'MATCH(${P}{FIRST},$A${FIRST}:$A${last},0))<>${R}{FIRST},FALSE))'], **orange()))
    for r, row in enumerate(rows, FIRST):
        for i, c in enumerate(cols, 1):
            v = row.get(c["name"])
            if v is not None:
                cell = ws.cell(r, i, v)
                if c["kind"] == "date":
                    cell.number_format = "yyyy-mm-dd"
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A2:{get_column_letter(len(cols))}{last}"


SETTINGS = [  # key, value, unit, meaning
    ("walkway_width", 60, "cm", "Narrowest gap that counts as a path from a door."),
    ("reach", 30, "cm", "A clear zone this close to a path counts as reachable."),
    ("person_height", 200, "cm", "Clear zones on the floor are kept free up to this height."),
    ("blocks_walking_below", 150, "cm", "Anything whose solid part starts lower than this is in the way when walking."),
    ("circuit_limit", 80, "%", "Running load allowed on a circuit, as a share of its breaker rating."),
    ("heavy_load", 1000, "W", "Peak load that puts a critical device on the same circuit at risk."),
    ("grid", 5, "cm", "Raster size for walkways and bench space. Smaller is finer but slower."),
    ("expiry_warning_days", 30, "days", "Warn this long before a document on the documents sheet expires."),
    ("fit_margin", 2, "cm", "Slack for measuring error and ventilation: under a bench, below the ceiling, through a door."),
    ("door_gap", 10, "cm", "Kept free beside a door's hinge (door column) so it opens past 90 degrees."),
    ("utility_reach", 300, "cm", "How far a gas, water, drain, vacuum, air, network or exhaust point may be (needs column)."),
    ("sprinkler_clearance", 45, "cm", "In rooms with sprinklers, nothing may reach higher than this below the ceiling."),
    ("sash_clearance", 15, "cm", "Work inside a fume hood is kept at least this far behind the sash."),
]


def write_settings(wb):
    ws = wb.create_sheet("settings")
    ws.sheet_properties.tabColor = "595959"
    for c, (text, dark) in enumerate((("setting", True), ("value", True), ("unit", False), ("what it means", False)), 1):
        header(ws.cell(1, c), text, dark)
    for r, row in enumerate(SETTINGS, 2):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    dv = DataValidation(type="whole", operator="greaterThan", formula1="0", errorTitle="Whole number",
                        error="A whole number above 0, in the unit shown.")
    dv.add(f"B2:B{len(SETTINGS) + 1}")
    ws.add_data_validation(dv)
    for col, width in zip("ABCD", (22, 9, 6, 80)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"


def write_lists(wb, extra):
    ws = wb.create_sheet("lists")
    ws.sheet_properties.tabColor = "A5A5A5"
    col = 1
    for name, values in lists.LISTS.items():
        values = values + extra.get(name, [])
        a, b = get_column_letter(col), get_column_letter(col + 1)
        header(ws.cell(1, col), name, True)
        ws.cell(1, col).comment = Comment("Dropdown source. To add a value, INSERT a row above the last entry "
                                          "so the dropdown range grows with it.", "lab-map", width=260, height=80)
        header(ws.cell(1, col + 1), lists.SECOND_HEADER.get(name, "meaning"), False)
        for r, (v, m) in enumerate(values, 2):
            ws.cell(r, col, v)
            ws.cell(r, col + 1, m if m != "" else None)
        wb.defined_names[f"L_{name}"] = DefinedName(f"L_{name}", attr_text=f"lists!${a}$2:${a}${len(values) + 1}")
        ws.column_dimensions[a].width = max(9, max(len(str(v)) for v, _ in values) + 2)
        ws.column_dimensions[b].width = max(10, min(38, max(len(str(m)) for _, m in values) + 2))
        ws.column_dimensions[get_column_letter(col + 2)].width = 2
        col += 3
    header(ws.cell(1, col), "ID prefix", True)
    header(ws.cell(1, col + 1), "used for", False)
    for r, (p, m) in enumerate(lists.PREFIXES, 2):
        ws.cell(r, col, p)
        ws.cell(r, col + 1, m)
    ws.column_dimensions[get_column_letter(col)].width = 13
    ws.column_dimensions[get_column_letter(col + 1)].width = 30
    ws.freeze_panes = "A2"


README = [
    ("Lab map: data workbook", "title"),
    ("Row 1 is the column name: hover over it for an explanation. Row 2 is a short hint. Data starts on row 3.", ""),
    ("", ""),
    ("Colours", "bold"),
    ("required", "dark"),
    ("optional", "light"),
    ("Red cell: duplicate ID, or an ID that doesn't exist on the sheet it points to.", ""),
    ("Orange room cell: differs from its parent's room.", ""),
    ("", ""),
    ("Suggested order: one walk-around each", "bold"),
    ("0. Triage: list every piece of equipment (id, name, category, room on placeables) and fill plan / usage / "
     "condition on equipment. Decide what leaves BEFORE you spend time measuring it.", ""),
    ("1. rooms: measure each room and draw its outline in rooms/<ID>.svg (copy rooms/_TEMPLATE.svg).", ""),
    ("2. placeables: tape measure. Dimensions and clearances. Positions are optional: you can drag things into "
     "place in Inkscape later.", ""),
    ("3. equipment: nameplates and people. Power, owner, usage.", ""),
    ("4. circuits + services: the distribution board, then every socket, strip and tap.", ""),
    ("5. links: which PC drives which instrument; cable and tubing limits.", ""),
    ("6. items: drawer contents. Start with the ~50 things people hunt for, then grow on demand. Spare parts go "
     "here too: spare_for says which equipment, min_qty how many to keep, rs_part / buy_link where to order.", ""),
    ("", ""),
    ("Conventions", "bold"),
    ("Units: centimetres, watts, kilograms. Whole numbers only: round to the nearest cm.", ""),
    ("IDs: UPPERCASE letters, digits, - and . only. Never put the location in an ID (things move). "
     "Never reuse one.", ""),
    ("Positions are in the parent's frame. In a room: x from the drawing's left edge, y from its top edge. "
     "On a bench: x from its left end as you face it, y from its back edge.", ""),
    ("faces = which way the front points: N, NE, E, SE, S, SW, W or NW. N = top of the drawing, not compass north. "
     "For things on/in/under a parent, blank = same way as the parent.", ""),
    ("Blank x and y = not placed yet: it appears in a staging area beside the room drawing, ready to drag in.", ""),
    ("Tables, desks and open benches: free_under = clear height underneath. Blank = solid to the floor.", ""),
    ("Stacking: anything can sit on anything (mount = on). Mark the lower one stackable = yes, or you get a warning.", ""),
    ("settings sheet: the thresholds the checks use (walkway width, circuit limit, ...). Change them to match your rules.", ""),
    ("keep_apart sheet: which tags (tags column on placeables) must be kept how far apart, e.g. vibrates / "
     "vibration-sensitive. Add your own rows.", ""),
    ("Photos and SOPs aren't listed here: they're linked by file name (see README.md).", ""),
    ("", ""),
]


def write_readme(ws, example):
    ws.column_dimensions["A"].width = 125
    ws.sheet_view.showGridLines = False
    last = ("This is the worked EXAMPLE. Your own data goes in ../lab-data.xlsx." if example else
            "Full guide: README.md in this folder. Worked example: example/lab-data.xlsx.")
    for r, (text, style) in enumerate(README + [(last, "bold")], 1):
        c = ws.cell(r, 1, text or None)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if style == "title":
            c.font = Font(bold=True, size=16, color=DARK)
        elif style == "bold":
            c.font = Font(bold=True, size=12)
        elif style in ("dark", "light"):
            header(c, text, style == "dark")


def build(path, example):
    wb = Workbook()
    wb.active.title = "_readme"
    write_readme(wb.active, example)
    for name, tab, freeze, nrows, cols, _ in schema.SHEETS:
        rows = ex.ROWS.get(name, schema.DEFAULT_ROWS.get(name, [])) if example else schema.DEFAULT_ROWS.get(name, [])
        write_sheet(wb, name, tab, freeze, nrows, cols, rows)
    write_settings(wb)
    write_lists(wb, ex.EXTRA_LISTS if example else {})
    wb.properties.title = "Lab map data"
    wb.properties.creator = "lab-map template"
    wb.save(path)
    print("wrote", path)


if __name__ == "__main__":
    # Regenerates the empty template and the example (and the copy example/reset.bat restores from).
    # Run from anywhere:  python tools/make_workbooks.py [repository folder]
    import shutil
    from pathlib import Path

    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent
    build(out / "lab-data.xlsx", False)
    build(out / "example" / "lab-data.xlsx", True)
    shutil.copy2(out / "example" / "lab-data.xlsx", out / "example" / "lab-data.original.xlsx")
