"""Load a lab folder: lab-data.xlsx plus the SVG room outlines and shape profiles.

Problems with the data itself (unknown IDs, contradictions, values that aren't numbers) become findings with
rule "data". Data that's merely missing isn't a problem: the report counts it under completeness instead.
"""
from __future__ import annotations

import datetime as dt
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path

from . import svg

SHEETS = ("rooms", "placeables", "equipment", "services", "circuits", "links", "items", "documents", "keep_apart")
ID_SHEETS = ("rooms", "placeables", "equipment", "services", "circuits", "items", "documents")
DATES = {("placeables", "checked"), ("placeables", "decommissioned"), ("documents", "filled"), ("documents", "expires")}
ID_COLUMNS = {"id", "room", "parent", "outlet", "circuit", "fed_by", "container", "from", "to"}
NUMERIC = {
    "rooms": {"width", "depth", "ceiling", "cooling"},
    "placeables": {"x", "y", "z", "w", "d", "h", "free_under", "clear_front", "clear_back", "clear_left",
                   "clear_right", "clear_top", "fill", "inner_w", "inner_d", "inner_h", "inner_z"},
    "equipment": {"plugs", "watts_typ", "watts_max", "volts", "weight"},
    "services": {"x", "y", "z", "sockets", "rating_a"},
    "circuits": {"rating_a", "volts", "phase"},
    "links": {"max_len"},
    "items": {"min_qty"},
    "documents": set(),
    "keep_apart": {"distance"},
}
TEXT = {("rooms", "floor"), ("equipment", "serial"), ("equipment", "asset_tag"), ("items", "qty"), ("items", "rs_part")}
REQUIRED = {"rooms": ("id",), "placeables": ("id", "room", "mount"), "equipment": ("id",),
            "services": ("id", "type", "room"), "circuits": ("id",), "links": ("from", "to", "type"),
            "items": ("id",), "documents": ("id", "applies_to", "type", "status"),
            "keep_apart": ("tag", "away_from", "distance")}
REFS = [  # (sheet, column, sheet the value must be on)
    ("placeables", "room", "rooms"), ("placeables", "parent", "placeables"),
    ("equipment", "id", "placeables"), ("equipment", "outlet", "services"),
    ("services", "room", "rooms"), ("services", "parent", "placeables"), ("services", "circuit", "circuits"),
    ("services", "fed_by", "services"), ("links", "from", "placeables"), ("links", "to", "placeables"),
    ("items", "container", "placeables"),
]
ENUMS = {  # (sheet, column) -> list on the lists sheet
    ("placeables", "category"): "category", ("placeables", "mount"): "mount", ("placeables", "faces"): "faces",
    ("placeables", "fixed"): "yesno", ("placeables", "mobile"): "yesno", ("placeables", "stackable"): "yesno",
    ("placeables", "door"): "door", ("rooms", "sprinklers"): "yesno", ("keep_apart", "level"): "level",
    ("equipment", "condition"): "condition", ("equipment", "plan"): "plan", ("equipment", "usage"): "usage",
    ("equipment", "usage_source"): "usage_source", ("equipment", "critical"): "yesno",
    ("services", "type"): "service_type", ("circuits", "rcd"): "yesno", ("circuits", "backed"): "backed",
    ("circuits", "phase"): "phase", ("links", "type"): "link_type", ("items", "category"): "item_category",
    ("documents", "type"): "doc_type", ("documents", "status"): "doc_status",
}
KNOWN_LISTS = set(ENUMS.values()) | {"shape", "plug_type"}
FALLBACK_LISTS = {
    "mount": ["floor", "wall", "on", "under", "in", "part"],
    "faces": ["N", "NE", "E", "SE", "S", "SW", "W", "NW"],
    "yesno": ["yes", "no"],
    "service_type": ["outlet", "strip", "gas", "vacuum", "air", "water", "drain", "network", "exhaust"],
    "door": ["left", "right", "both"],
    "level": ["problem", "warning"],
}
UTILITIES = ("gas", "vacuum", "air", "water", "drain", "network", "exhaust")  # what the needs column can ask for
FALLBACK_LINK_LEN = {"usb": 500, "usb3": 300, "ethernet": 10000, "serial": 1500, "gpib": 200, "video": 500}
SETTINGS = {  # the settings sheet can change these
    "walkway_width": 60, "reach": 30, "person_height": 200, "blocks_walking_below": 150,
    "circuit_limit": 80, "heavy_load": 1000, "grid": 5, "expiry_warning_days": 30,
    "fit_margin": 2, "door_gap": 10, "utility_reach": 300, "sprinkler_clearance": 45, "sash_clearance": 15,
}


@dataclass
class Finding:
    rule: str
    message: str
    ids: tuple = ()
    room: str | None = None


@dataclass
class Lab:
    folder: Path
    rooms: dict = field(default_factory=dict)
    placeables: dict = field(default_factory=dict)
    equipment: dict = field(default_factory=dict)
    services: dict = field(default_factory=dict)
    circuits: dict = field(default_factory=dict)
    links: list = field(default_factory=list)
    items: dict = field(default_factory=dict)
    documents: dict = field(default_factory=dict)
    keep_apart: list = field(default_factory=list)
    decommissioned: dict = field(default_factory=dict)  # id -> date it left (None if only plan = decommissioned)
    gone: set = field(default_factory=set)  # decommissioned, with their drawers and parts
    lists: dict = field(default_factory=dict)
    settings: dict = field(default_factory=lambda: dict(SETTINGS))
    issues: list = field(default_factory=list)
    _profiles: dict = field(default_factory=dict, repr=False)
    _children: dict | None = field(default=None, repr=False)

    def issue(self, message, ids=(), room=None):
        self.issues.append(Finding("data", message, tuple(ids), room))

    def allowed(self, list_name):
        vals = [v for v, _ in self.lists[list_name]] if list_name in self.lists else FALLBACK_LISTS.get(list_name)
        if vals is None:
            return None
        return {v.upper() if list_name == "faces" else v.lower() if isinstance(v, str) else v for v in vals}

    def link_default(self, kind):
        table = {str(v).lower(): m for v, m in self.lists.get("link_type", [])} or FALLBACK_LINK_LEN
        m = table.get(kind)
        return m if isinstance(m, (int, float)) else None

    @property
    def children(self):
        if self._children is None:
            self._children = {}
            for i, r in self.placeables.items():
                if r.get("parent"):
                    self._children.setdefault(r["parent"], []).append(i)
        return self._children

    def profile(self, name):
        """(footprint, clearance or None) from shapes/<name>.svg, or None if it can't be read."""
        if name not in self._profiles:
            path, got = self.folder / "shapes" / f"{name}.svg", None
            if path.exists():
                try:
                    shapes = svg.outlines(path)
                    if len(shapes.get("footprint", [])) >= 3:
                        got = (shapes["footprint"], shapes.get("clearance"))
                except (ET.ParseError, ValueError, IndexError):
                    got = None
            self._profiles[name] = got
        return self._profiles[name]


def load(folder, moves=None):
    """A Lab from a folder. moves: {sheet: {id: {column: value}}} applied on top of the workbook, to try an
    arrangement from the Inkscape layout without writing it (see layout.moves_from)."""
    folder = Path(folder)
    rows, lists = read_workbook(folder / "lab-data.xlsx")
    for sheet, changes in (moves or {}).items():
        for r in rows.get(sheet, []):
            i = str(r.get("id") or "").strip().upper()
            if i in changes:
                r.update(changes[i])
    return build(folder, rows, lists, settings=rows.pop("_settings", None))


def read_workbook(path):
    """({sheet: rows, "_settings": {key: value}}, lists)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = {name: _read_sheet(wb[name]) if name in wb.sheetnames else [] for name in SHEETS}
        lists = _read_lists(wb["lists"]) if "lists" in wb.sheetnames else {}
        if "settings" in wb.sheetnames:
            it = wb["settings"].iter_rows(values_only=True)
            next(it, None)
            rows["_settings"] = {str(r[0]).strip(): _clean(r[1]) for r in it if r and r[0] and len(r) > 1}
    finally:
        wb.close()
    _fill_formulas(path, rows)
    return rows, lists


def _fill_formulas(path, rows):
    """Formula cells whose stored result is missing (e.g. after a save by openpyxl, as `pull` does): work out
    simple arithmetic like =230*2.5 ourselves; anything else is left blank with a note in the row."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        for name, recs in rows.items():
            if name not in wb.sheetnames:
                continue
            it = wb[name].iter_rows(values_only=True)
            header = [h.strip() if isinstance(h, str) else None for h in next(it, ())]
            by_row = {r["_row"]: r for r in recs}
            for n, values in enumerate(it, start=2):
                rec = by_row.get(n)
                for col, v in zip(header, values):
                    if rec is None or not col or not isinstance(v, str) or not v.startswith("=") or rec.get(col) is not None:
                        continue
                    value = _arithmetic(v[1:])
                    if value is None:
                        rec.setdefault("_formulas", []).append(col)
                    else:
                        rec[col] = int(value) if float(value).is_integer() else value
    finally:
        wb.close()


def _arithmetic(expr):
    """Value of an expression made only of numbers, + - * / and brackets; None for anything else."""
    import ast

    def ev(node):
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return node.value
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.USub, ast.UAdd)):
            v = ev(node.operand)
            return -v if isinstance(node.op, ast.USub) else v
        if isinstance(node, ast.BinOp) and isinstance(node.op, (ast.Add, ast.Sub, ast.Mult, ast.Div)):
            a, b = ev(node.left), ev(node.right)
            if isinstance(node.op, ast.Add):
                return a + b
            if isinstance(node.op, ast.Sub):
                return a - b
            if isinstance(node.op, ast.Mult):
                return a * b
            return a / b
        raise ValueError

    try:
        return ev(ast.parse(expr, mode="eval").body)
    except (ValueError, SyntaxError, ZeroDivisionError):
        return None


def write_moves(path, moves, backup_dir):
    """Write {sheet: {id: {column: value}}} into the workbook after copying it to backup_dir. Only those cells
    change; validation, formatting and comments are kept. Returns (backup path, ids not found)."""
    import shutil
    from openpyxl import load_workbook

    path, backup_dir = Path(path), Path(backup_dir)
    wb = load_workbook(path)
    missing = []
    for sheet, changes in moves.items():
        if not changes:
            continue
        ws = wb[sheet]
        col = {c.value.strip(): c.column for c in ws[1] if isinstance(c.value, str)}
        rows = {}
        for n in range(3, ws.max_row + 1):
            v = ws.cell(n, col["id"]).value
            if v is not None:
                rows.setdefault(str(v).strip().upper(), n)
        for i, change in changes.items():
            n = rows.get(i)
            if n is None:
                missing.append(i)
                continue
            for k, v in change.items():
                if k in col:
                    ws.cell(n, col[k]).value = v
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = backup_dir / f"lab-data-{dt.datetime.now():%Y%m%d-%H%M%S}.xlsx"
    shutil.copy2(path, backup)
    wb.save(path)
    return backup, missing


def _clean(v):
    if isinstance(v, str):
        return v.strip() or None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, dt.datetime):
        return v.date()
    return v


def _read_sheet(ws):
    it = ws.iter_rows(values_only=True)
    header = [h.strip() if isinstance(h, str) else None for h in next(it, ())]
    next(it, None)  # row 2 holds the hints
    out = []
    for n, values in enumerate(it, start=3):
        rec = {c: _clean(v) for c, v in zip(header, values) if c}
        if any(v is not None for v in rec.values()):
            rec["_row"] = n
            out.append(rec)
    return out


def _read_lists(ws):
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    lists = {}
    for c, name in enumerate(grid[0] if grid else []):
        if not isinstance(name, str) or name.strip() not in KNOWN_LISTS:
            continue
        vals = []
        for row in grid[1:]:
            v = _clean(row[c]) if c < len(row) else None
            if v is None:
                break
            vals.append((v, _clean(row[c + 1]) if c + 1 < len(row) else None))
        lists[name.strip()] = vals
    return lists


def label(sheet, r):
    if sheet == "links":
        return f"link {r.get('from')} → {r.get('to')}"
    return r.get("id") or f"{sheet} row {r.get('_row', '?')}"


def _normalise(lab, sheet, r):
    where = f"{sheet} row {r.get('_row', '?')}"
    for col in r.pop("_formulas", []):
        lab.issue(f"{where}: the formula in {col} has no stored result. Open lab-data.xlsx in Excel and save it once",
                  [r["id"]] if r.get("id") else [])
    for k, v in list(r.items()):
        if k == "_row" or v is None:
            continue
        if (sheet, k) in TEXT:
            r[k] = str(v)
        elif k in NUMERIC[sheet]:
            if isinstance(v, str):
                try:
                    v = float(v.replace(",", "."))
                except ValueError:
                    lab.issue(f"{where}: {k} = {v!r} isn't a number", [r["id"]] if r.get("id") else [])
                    r[k] = None
                    continue
            if isinstance(v, bool) or not isinstance(v, (int, float)):
                lab.issue(f"{where}: {k} = {v!r} isn't a number")
                r[k] = None
                continue
            r[k] = int(v) if float(v).is_integer() else v
        elif k in ID_COLUMNS:
            r[k] = str(v).strip().upper()
        elif (sheet, k) in DATES and not isinstance(v, dt.date):
            try:
                r[k] = dt.date.fromisoformat(str(v).strip()[:10])
            except ValueError:
                lab.issue(f"{where}: {k} = {v!r} isn't a date (use YYYY-MM-DD)", [r["id"]] if r.get("id") else [])
                r[k] = None
        elif (sheet, k) in (("documents", "applies_to"), ("items", "spare_for")):
            r[k] = [s.strip().upper() for s in str(v).replace(",", ";").split(";") if s.strip()]
        elif (sheet, k) in (("placeables", "tags"),):
            r[k] = [s.strip().lower() for s in str(v).replace(",", ";").split(";") if s.strip()]
        elif sheet == "keep_apart" and k in ("tag", "away_from"):
            r[k] = str(v).strip().lower()
        elif (sheet, k) == ("equipment", "needs"):
            got = []
            for s in str(v).replace(",", ";").split(";"):
                kind, _, medium = s.strip().partition(":")
                kind = kind.strip().lower()
                if not kind:
                    continue
                if kind not in UTILITIES:
                    lab.issue(f"{label(sheet, r)}: needs {s.strip()!r}: use {', '.join(UTILITIES)} (optionally :medium)",
                              [r.get("id")] if r.get("id") else [])
                    continue
                got.append((kind, medium.strip() or None))
            r[k] = got
    for (sh, col), list_name in ENUMS.items():
        v = r.get(col)
        allowed = lab.allowed(list_name) if sh == sheet and v is not None else None
        if allowed is None:
            continue
        key = str(v).upper() if col == "faces" else v.lower() if isinstance(v, str) else v
        if key in allowed:
            r[col] = key
        else:
            lab.issue(f"{label(sheet, r)}: {col} = {v!r} isn't on the {list_name} list", [r.get("id")] if r.get("id") else [])
    shape = r.get("shape") if sheet == "placeables" else None
    if shape and shape not in ("rect", "circle", "group") and not str(shape).startswith("@"):
        lab.issue(f"{label(sheet, r)}: shape = {shape!r}: use rect, circle, group or @profile-name", [r["id"]])


def build(folder, rows, lists=None, room_polys=None, settings=None):
    """A Lab from sheet rows ({sheet: [ {column: value} ]}). room_polys can stand in for the room SVGs (tests)."""
    lab = Lab(folder=Path(folder), lists=lists or {})
    for key, value in (settings or {}).items():
        if key not in SETTINGS:
            lab.issue(f"settings: '{key}' isn't a setting the checks know")
        elif value is None:
            continue
        elif isinstance(value, (int, float)) and not isinstance(value, bool) and value > 0:
            lab.settings[key] = value
        else:
            lab.issue(f"settings: {key} = {value!r} should be a number above 0; using {SETTINGS[key]}")
    tables = {}
    for sheet in SHEETS:
        recs = []
        for r in rows.get(sheet, []):
            r = dict(r)
            _normalise(lab, sheet, r)
            missing = [k for k in REQUIRED[sheet] if r.get(k) is None]
            if missing:
                lab.issue(f"{sheet} row {r.get('_row', '?')}: no {', '.join(missing)}", [r["id"]] if r.get("id") else [])
                if "id" in missing or "from" in missing or "to" in missing:
                    continue
            recs.append(r)
        tables[sheet] = recs
    for sheet in ID_SHEETS:
        table = {}
        for r in tables[sheet]:
            if r["id"] in table:
                lab.issue(f"{sheet}: {r['id']} is used twice (rows {table[r['id']].get('_row')} and {r.get('_row')})",
                          [r["id"]])
            else:
                table[r["id"]] = r
        setattr(lab, sheet, table)
    lab.links = tables["links"]
    lab.keep_apart = tables["keep_apart"]
    for sheet, col, target in REFS:
        pool = getattr(lab, target)
        for r in lab.links if sheet == "links" else getattr(lab, sheet).values():
            v = r.get(col)
            if v is not None and v not in pool:
                lab.issue(f"{label(sheet, r)}: {col} {v} isn't on the {target} sheet", [r.get("id") or v], r.get("room"))
    for d in lab.documents.values():
        for i in d.get("applies_to") or []:
            if i not in lab.placeables and i not in lab.services and i not in lab.rooms:
                lab.issue(f"{d['id']}: applies_to {i} isn't a placeable, socket or room", [d["id"]])
    for it in lab.items.values():
        if not it.get("container") and not it.get("elsewhere"):
            lab.issue(f"{it['id']}: no container, and nothing under elsewhere either", [it["id"]])
        for i in it.get("spare_for") or []:
            if i not in lab.placeables:
                lab.issue(f"{it['id']}: spare_for {i} isn't on the placeables sheet", [it["id"]])
    _decommissioned(lab)
    _check_placeables(lab)
    _check_services(lab)
    _attach_rooms(lab, room_polys)
    return lab


RS_URL = "https://uk.rs-online.com/web/c/?searchTerm={}"  # RS search: a stock number goes straight to the product


def stock(qty):
    """How many there are, from the free-text qty: the first number in it (3, ~20, 2 boxes). None if there's none."""
    m = re.search(r"\d+(?:[.,]\d+)?", str(qty if qty is not None else ""))
    return float(m.group().replace(",", ".")) if m else None


def order_links(item):
    """[(label, url)] for ordering an item: RS stock number first, then its buy_link."""
    out = []
    if item.get("rs_part"):
        out.append((f"RS {item['rs_part']}", RS_URL.format(re.sub(r"[^0-9A-Za-z]", "", item["rs_part"]))))
    link = str(item.get("buy_link") or "").strip()
    if link:
        host = re.match(r"https?://(?:www\.)?([^/]+)", link)
        out.append((host.group(1) if host else "buy link", link))
    return out


def _decommissioned(lab):
    """Which rows are gone: a decommissioned date on the placeables sheet, or plan = decommissioned. What's part of
    them (drawers, cabinet shelves, the pieces of a composed bench) goes with them."""
    P = lab.placeables
    for i, r in P.items():
        if r.get("decommissioned") or lab.equipment.get(i, {}).get("plan") == "decommissioned":
            lab.decommissioned[i] = r.get("decommissioned")
    lab.gone = set(lab.decommissioned)
    grew = True
    while grew:
        grew = False
        for i, r in P.items():
            if i not in lab.gone and r.get("parent") in lab.gone and (
                    r.get("mount") == "part" or (r.get("mount") == "in" and r.get("x") is None)):
                lab.gone.add(i)
                grew = True


def _check_placeables(lab):
    P = lab.placeables
    for i, r in P.items():
        mount, parent = r.get("mount"), r.get("parent")
        pr = P.get(parent)
        if pr and pr.get("room") and r.get("room") and pr["room"] != r["room"]:
            lab.issue(f"{i} is in {r['room']} but its parent {parent} is in {pr['room']}", [i, parent], r["room"])
        if mount in ("on", "under", "in", "part") and not parent and (mount in ("in", "part") or r.get("x") is not None):
            lab.issue(f"{i} is mounted '{mount}' but has no parent", [i], r.get("room"))
        if mount in ("floor", "wall") and parent:
            lab.issue(f"{i} has parent {parent} but mount '{mount}': use on, under, in or part, or clear the parent",
                      [i], r.get("room"))
        if mount == "in" and r.get("x") is not None and pr and not (pr.get("inner_w") and pr.get("inner_d")):
            lab.issue(f"{i} is placed inside {parent}, but {parent} has no inner_w and inner_d (its working space)",
                      [i, parent], r.get("room"))
        if mount == "part" and pr and pr.get("shape") != "group":
            lab.issue(f"{i} is a 'part' but its parent {parent} isn't shape = group", [i], r.get("room"))
        if r.get("shape") == "group" and (parent or mount != "floor"):
            lab.issue(f"{i}: groups stand on the room floor (mount = floor, no parent)", [i], r.get("room"))
        if (r.get("x") is None) != (r.get("y") is None):
            lab.issue(f"{i} has only one of x and y", [i], r.get("room"))
        if mount == "wall" and r.get("x") is not None and r.get("z") is None:
            lab.issue(f"{i} is wall-mounted but has no z (height of its underside)", [i], r.get("room"))
        seen, j = {i}, parent
        while j in P:
            if j in seen:
                lab.issue(f"{i}: its parent chain loops back on itself", [i], r.get("room"))
                break
            seen.add(j)
            j = P[j].get("parent")
    for e in lab.equipment.values():
        s, p = lab.services.get(e.get("outlet")), P.get(e["id"])
        if s and p and s.get("room") != p.get("room"):
            lab.issue(f"{e['id']} is in {p.get('room')} but its outlet {s['id']} is in {s.get('room')}",
                      [e["id"], s["id"]], p.get("room"))


def _check_services(lab):
    for i, s in lab.services.items():
        p = lab.placeables.get(s.get("parent"))
        if p and p.get("room") != s.get("room"):
            lab.issue(f"{i} is in {s.get('room')} but its parent {p['id']} is in {p.get('room')}", [i], s.get("room"))
        seen, j = {i}, s.get("fed_by")
        while j in lab.services:
            if j in seen:
                lab.issue(f"{i}: its fed_by chain loops back on itself", [i], s.get("room"))
                break
            seen.add(j)
            j = lab.services[j].get("fed_by")


def _attach_rooms(lab, room_polys):
    for rid, room in lab.rooms.items():
        if room_polys is not None:
            room["poly"] = room_polys.get(rid)
            continue
        shell = room.get("shell")
        if not shell:
            continue  # counted under completeness
        path = lab.folder / "rooms" / shell
        if not path.exists():
            lab.issue(f"{rid}: rooms/{shell} doesn't exist", [rid], rid)
            continue
        try:
            poly = svg.outlines(path).get("interior")
        except (ET.ParseError, ValueError, IndexError) as e:
            lab.issue(f"{rid}: can't read rooms/{shell} ({e})", [rid], rid)
            continue
        if not poly or len(poly) < 3:
            lab.issue(f"{rid}: rooms/{shell} has no outline with id 'interior'", [rid], rid)
            continue
        room["poly"] = poly
        xs, ys = [p[0] for p in poly], [p[1] for p in poly]
        if abs(min(xs)) > 1 or abs(min(ys)) > 1:
            lab.issue(f"{rid}: the outline's top-left is at {min(xs):.0f},{min(ys):.0f}, not 0,0", [rid], rid)
        size = (max(xs) - min(xs), max(ys) - min(ys))
        want = (room.get("width"), room.get("depth"))
        if all(want) and (abs(size[0] - want[0]) > 2 or abs(size[1] - want[1]) > 2):
            lab.issue(f"{rid}: rooms sheet says {want[0]} × {want[1]} cm, the outline is {size[0]:.0f} × {size[1]:.0f}",
                      [rid], rid)
